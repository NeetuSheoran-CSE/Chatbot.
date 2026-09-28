from __future__ import annotations

import io
import os
import sys
import threading
from pathlib import Path
from typing import List, Optional

import requests
import uvicorn
from fastapi import FastAPI, File, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

# Ensure UTF-8 output handling for Windows terminal
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# Set threading and OpenBLAS environment limits early to prevent BLAS memory issues on Windows
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["OMP_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"


# ---------------------------------------------------------------------------
# Environment Variable Loading
# ---------------------------------------------------------------------------
def load_env():
    """Load KEY=value lines from .env with UTF-8 BOM tolerance."""
    env_file = Path(__file__).parent / ".env"
    if env_file.exists():
        for line in env_file.read_text(encoding="utf-8-sig").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                key, value = line.split("=", 1)
                os.environ.setdefault(
                    key.strip(),
                    value.strip().strip('"').strip("'")
                )


load_env()

# Optional dotenv fallback
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

API_URL = "https://api.groq.com/openai/v1/chat/completions"
MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")
EMBEDDING_MODEL_NAME = os.getenv("EMBEDDING_MODEL_NAME", "sentence-transformers/all-MiniLM-L6-v2")

SYSTEM_PROMPT = {
    "role": "system",
    "content": (
        "You are a friendly, clear assistant. "
        "Keep answers concise and use Markdown when it helps."
    ),
}

# Document parsers, embeddings and the vector store are imported lazily inside
# the functions that need them: importing them at startup delayed port binding
# and exceeded Render's 512MB memory limit.


# ---------------------------------------------------------------------------
# Pydantic Schemas
# ---------------------------------------------------------------------------
class Message(BaseModel):
    role: str
    content: str


class ChatRequest(BaseModel):
    messages: list[Message]


class QueryRequest(BaseModel):
    question: str = Field(..., min_length=1, description="Question to ask.")
    top_k: int = Field(default=4, ge=1, le=10, description="Number of relevant chunks to retrieve.")


class SourceChunk(BaseModel):
    content: str
    metadata: dict


class QueryResponse(BaseModel):
    question: str
    answer: str
    reply: Optional[str] = None
    sources: List[SourceChunk] = []


class UploadResponse(BaseModel):
    success: bool
    filename: str
    total_chunks: int
    message: str


class StatusResponse(BaseModel):
    document_loaded: bool
    filename: Optional[str] = None
    total_chunks: int = 0
    embedding_model: str
    llm_model: str


# ---------------------------------------------------------------------------
# Direct Groq API Client (requests-based)
# ---------------------------------------------------------------------------
def call_groq_api(messages: list[dict], model: Optional[str] = None, temperature: Optional[float] = None) -> str:
    """Direct HTTP request to Groq API as defined in the user's reference code."""
    api_key = os.environ.get("GROQ_API_KEY", "").strip()
    if not api_key:
        raise HTTPException(
            status_code=500,
            detail="GROQ_API_KEY is not set on the server. Please check your .env file."
        )

    target_model = model or os.environ.get("GROQ_MODEL", MODEL)
    payload = {
        "model": target_model,
        "messages": messages
    }
    if temperature is not None:
        payload["temperature"] = temperature

    try:
        response = requests.post(
            API_URL,
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json"
            },
            json=payload,
            timeout=60
        )
    except requests.RequestException as e:
        raise HTTPException(
            status_code=502,
            detail=f"Could not reach Groq API: {str(e)}"
        )

    if response.status_code != 200:
        print("Groq error:", response.text)
        raise HTTPException(
            status_code=502,
            detail=f"Groq API returned {response.status_code}: {response.text}"
        )

    try:
        data = response.json()
        reply = (data["choices"][0]["message"].get("content") or "").strip()
    except (KeyError, IndexError, ValueError):
        raise HTTPException(
            status_code=502,
            detail="Unexpected response from Groq API."
        )

    if not reply:
        raise HTTPException(
            status_code=502,
            detail="The AI returned an empty response."
        )

    return reply


# ---------------------------------------------------------------------------
# Document Extractors
# ---------------------------------------------------------------------------
def extract_text_from_pdf(file_bytes: bytes, filename: str) -> List[dict]:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    docs: List[dict] = []
    for page_idx, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        text = text.strip()
        if text:
            docs.append({
                "content": text,
                "metadata": {"source": filename, "page": page_idx, "file_type": "pdf"}
            })
    if not docs:
        raise ValueError("Could not extract any readable text from the uploaded PDF.")
    return docs


def extract_text_from_docx(file_bytes: bytes, filename: str) -> List[dict]:
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted_paragraphs: List[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            extracted_paragraphs.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                extracted_paragraphs.append(" | ".join(row_cells))
    full_text = "\n\n".join(extracted_paragraphs).strip()
    if not full_text:
        raise ValueError("Could not extract any readable text from the uploaded DOCX.")
    return [{"content": full_text, "metadata": {"source": filename, "file_type": "docx"}}]


def split_text(
    text: str,
    chunk_size: int = 800,
    chunk_overlap: int = 120,
    separators: tuple = ("\n\n", "\n", ". ", " ", ""),
) -> List[str]:
    text = text.strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]
    if not separators:
        return [text[i:i + chunk_size] for i in range(0, len(text), chunk_size)]
    sep, rest = separators[0], separators[1:]
    if sep and sep not in text:
        return split_text(text, chunk_size, chunk_overlap, rest)
    pieces = text.split(sep) if sep else [text[i:i + chunk_size] for i in range(0, len(text), chunk_size)]

    chunks: List[str] = []
    current = ""
    for piece in pieces:
        piece = piece.strip()
        if not piece:
            continue
        if len(piece) > chunk_size:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(split_text(piece, chunk_size, chunk_overlap, rest))
            continue
        candidate = piece if not current else current + sep + piece
        if len(candidate) <= chunk_size:
            current = candidate
        else:
            chunks.append(current)
            current = current[-chunk_overlap:] + sep + piece if chunk_overlap else piece
            if len(current) > chunk_size:
                current = piece
    if current:
        chunks.append(current)
    return [c for c in chunks if c.strip()]


# ---------------------------------------------------------------------------
# RAG Service Manager
# ---------------------------------------------------------------------------
RAG_SYSTEM_PROMPT = (
    "You are an expert document analysis assistant. Your job is to answer user questions "
    "truthfully and concisely based on the provided context retrieved from the document.\n\n"
    "Guidelines:\n"
    "1. If the user greeting is conversational (like 'hello', 'hi', 'hlo', 'hey'), respond warmly and offer assistance with the document.\n"
    "2. Base factual answers strictly on the context provided below.\n"
    "3. If the answer cannot be found in the context, clearly state: "
    "'Based on the uploaded document, I cannot find information to answer this question.' Do not guess or fabricate information.\n"
    "4. Cite the relevant source chunks (e.g., [Source Chunk 1]) when referencing details.\n"
    "5. Format your output clearly with bullet points, paragraphs, or tables where appropriate."
)


class RAGManager:
    def __init__(self):
        self._lock = threading.Lock()
        self.chunks: List[dict] = []
        self._matrix = None
        self.current_filename: Optional[str] = None
        self.total_chunks: int = 0
        self._embedding_model = None

    @property
    def is_loaded(self) -> bool:
        return bool(self.chunks)

    @property
    def embedding_model(self):
        if self._embedding_model is None:
            from fastembed import TextEmbedding
            self._embedding_model = TextEmbedding(model_name=EMBEDDING_MODEL_NAME)
        return self._embedding_model

    def _embed(self, texts: List[str]):
        import numpy as np
        vectors = np.asarray(list(self.embedding_model.embed(texts)), dtype="float32")
        norms = np.linalg.norm(vectors, axis=1, keepdims=True)
        norms[norms == 0] = 1.0
        return vectors / norms

    def ingest_document(self, filename: str, file_bytes: bytes) -> int:
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            raw_docs = extract_text_from_pdf(file_bytes, filename)
        elif lower_name.endswith(".docx"):
            raw_docs = extract_text_from_docx(file_bytes, filename)
        else:
            raise ValueError("Unsupported file format. Please upload a PDF or DOCX file.")

        split_chunks: List[dict] = []
        for doc in raw_docs:
            for piece in split_text(doc["content"]):
                split_chunks.append({"content": piece, "metadata": dict(doc["metadata"])})
        if not split_chunks:
            raise ValueError("Document was parsed, but no text chunks could be produced.")

        for idx, chunk in enumerate(split_chunks, start=1):
            chunk["metadata"]["chunk_id"] = idx

        matrix = self._embed([chunk["content"] for chunk in split_chunks])
        with self._lock:
            self.chunks = split_chunks
            self._matrix = matrix
            self.current_filename = filename
            self.total_chunks = len(split_chunks)

        return self.total_chunks

    def answer_query(self, question: str, top_k: int = 4) -> dict:
        with self._lock:
            chunks, matrix = self.chunks, self._matrix
        if not chunks:
            # Fall back to direct Groq chat
            messages = [SYSTEM_PROMPT, {"role": "user", "content": question}]
            answer = call_groq_api(messages)
            return {"question": question, "answer": answer, "sources": []}

        import numpy as np
        similarities = matrix @ self._embed([question])[0]
        context_parts = []
        sources = []
        for idx in np.argsort(-similarities)[:top_k]:
            chunk = chunks[int(idx)]
            metadata = chunk["metadata"]
            page_info = f" (Page {metadata['page']})" if "page" in metadata else ""
            context_parts.append(
                f"[Source Chunk {metadata.get('chunk_id', '?')}{page_info}]:\n{chunk['content']}"
            )
            sources.append({
                "content": chunk["content"],
                "metadata": {
                    **metadata,
                    "relevance_distance": round(float(1 - similarities[idx]), 4)
                }
            })

        formatted_context = "\n\n---\n\n".join(context_parts)
        messages = [
            {"role": "system", "content": RAG_SYSTEM_PROMPT},
            {
                "role": "user",
                "content": f"Context from document:\n{formatted_context}\n\nQuestion:\n{question}\n\nAnswer:"
            },
        ]
        answer = call_groq_api(messages, temperature=0.2)

        return {
            "question": question,
            "answer": answer,
            "sources": sources
        }

    def clear(self):
        with self._lock:
            self.chunks = []
            self._matrix = None
            self.current_filename = None
            self.total_chunks = 0


rag_service = RAGManager()

# ---------------------------------------------------------------------------
# FastAPI Application & Middleware
# ---------------------------------------------------------------------------
app = FastAPI(
    title="Groq Chatbot",
    description="Groq LPU powered chatbot with both direct LLM chat and RAG document Q&A support.",
    version="1.1.0"
)

# Enable CORS for all origins to prevent 'Failed to fetch' browser errors
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------------------------------------------------------------------------
# Endpoints: Direct Chat API (User Reference Code Compatible)
# ---------------------------------------------------------------------------
@app.get("/api/health")
def health():
    return {
        "ok": bool(os.environ.get("GROQ_API_KEY")),
        "model": os.environ.get("GROQ_MODEL", MODEL)
    }


@app.post("/api/chat")
def chat(req: ChatRequest):
    history = [
        m.model_dump()
        for m in req.messages
        if m.role in ("user", "assistant")
    ]
    payload_messages = [SYSTEM_PROMPT] + history[-40:]
    reply = call_groq_api(payload_messages)
    return {"reply": reply}


# ---------------------------------------------------------------------------
# Endpoints: NexusRAG Frontend Support
# ---------------------------------------------------------------------------
@app.post("/ask", response_model=QueryResponse, summary="Ask questions (RAG or General Chat)")
async def ask_question(request: QueryRequest):
    try:
        if rag_service.is_loaded:
            result = rag_service.answer_query(request.question, top_k=request.top_k)
            return QueryResponse(
                question=result["question"],
                answer=result["answer"],
                reply=result["answer"],
                sources=[SourceChunk(content=s["content"], metadata=s["metadata"]) for s in result["sources"]]
            )
        else:
            # Free-form chat if no document is loaded
            messages = [SYSTEM_PROMPT, {"role": "user", "content": request.question}]
            reply = call_groq_api(messages)
            return QueryResponse(
                question=request.question,
                answer=reply,
                reply=reply,
                sources=[]
            )
    except HTTPException:
        raise
    except ValueError as ve:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(ve))
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Error generating answer: {str(e)}"
        )


@app.post("/upload", response_model=UploadResponse, summary="Upload and vectorize PDF/DOCX")
async def upload_document(file: UploadFile = File(...)):
    filename = file.filename or "unknown"
    lower_name = filename.lower()
    if not (lower_name.endswith(".pdf") or lower_name.endswith(".docx")):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Please upload a .pdf or .docx document."
        )

    try:
        content = await file.read()
        if len(content) == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded file is empty."
            )
        total_chunks = rag_service.ingest_document(filename, content)
        return UploadResponse(
            success=True,
            filename=filename,
            total_chunks=total_chunks,
            message=f"Successfully parsed '{filename}' and indexed {total_chunks} chunks."
        )
    except ValueError as ve:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(ve))
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"An error occurred while processing the document: {str(e)}"
        )


@app.get("/status", response_model=StatusResponse, summary="Get current RAG state")
async def get_status():
    return StatusResponse(
        document_loaded=rag_service.is_loaded,
        filename=rag_service.current_filename,
        total_chunks=rag_service.total_chunks,
        embedding_model=EMBEDDING_MODEL_NAME,
        llm_model=os.environ.get("GROQ_MODEL", MODEL)
    )


@app.post("/clear", summary="Clear current document and vector store")
async def clear_document():
    rag_service.clear()
    return {"message": "Document and vector store cleared successfully."}


# ---------------------------------------------------------------------------
# Static Web Assets Serving
# ---------------------------------------------------------------------------
static_dir = Path(__file__).parent / "static"
if static_dir.exists():
    app.mount("/static", StaticFiles(directory=str(static_dir)), name="static")


@app.get("/", include_in_schema=False)
def index():
    # Look for index.html in root or static/
    root_index = Path(__file__).parent / "index.html"
    static_index = static_dir / "index.html"
    if root_index.exists():
        return FileResponse(root_index)
    if static_index.exists():
        return FileResponse(static_index)
    return JSONResponse({"message": "Groq Chatbot API is running. index.html not found."})


# ---------------------------------------------------------------------------
# Server Launcher
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    if os.environ.get("GROQ_API_KEY"):
        print("[*] API key found.")
    else:
        print("[!] WARNING: GROQ_API_KEY not found.")

    port = int(os.environ.get("PORT", 8000))
    host = os.environ.get("HOST", "0.0.0.0")
    print(f"[*] Starting Groq Chatbot Server on http://{host}:{port}")
    uvicorn.run(app, host=host, port=port)